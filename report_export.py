"""Build an evidence-based Word report from the values shown in the analysis page."""

from datetime import datetime
from io import BytesIO
from math import ceil, isfinite
from pathlib import Path

import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt
from matplotlib.font_manager import FontProperties
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


_CHINESE_FONT_CANDIDATES = (
    r"C:\Windows\Fonts\msyh.ttc",
    r"C:\Windows\Fonts\simhei.ttf",
    r"C:\Windows\Fonts\simsun.ttc",
    "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
    "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
    "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
    "/System/Library/Fonts/PingFang.ttc",
    "/System/Library/Fonts/STHeiti Medium.ttc",
)


def _chart_font():
    for font_path in _CHINESE_FONT_CANDIDATES:
        path = Path(font_path)
        if path.is_file():
            return FontProperties(fname=str(path))
    return FontProperties(family="DejaVu Sans")


def _plot_style(font):
    return {
        "font.family": font.get_name(),
        "axes.unicode_minus": False,
        "axes.edgecolor": "#d1d5db",
        "axes.labelcolor": "#334155",
        "text.color": "#1f2937",
        "xtick.color": "#475569",
        "ytick.color": "#475569",
        "figure.facecolor": "white",
        "axes.facecolor": "white",
    }


def _as_list(trace, key):
    value = trace.get(key, [])
    if value is None:
        return []
    if hasattr(value, "tolist"):
        value = value.tolist()
    return list(value)


def _number(value):
    try:
        number = float(value)
    except (TypeError, ValueError):
        return 0.0
    return number if isfinite(number) else 0.0


def _chart_data(result, chart_name):
    traces = result.get("charts", {}).get(chart_name, {}).get("data", [])
    return traces[0] if traces else {}


def _to_png(figure):
    output = BytesIO()
    figure.savefig(
        output,
        format="png",
        dpi=150,
        bbox_inches="tight",
        facecolor="white",
    )
    plt.close(figure)
    output.seek(0)
    return output.getvalue()


def _build_trend_chart(result, font):
    trace = _chart_data(result, "sales_trend")
    dates = [str(value) for value in _as_list(trace, "x")]
    amounts = [_number(value) for value in _as_list(trace, "y")]

    with plt.rc_context(_plot_style(font)):
        figure, axis = plt.subplots(figsize=(9, 4.6))
        if dates and amounts:
            axis.plot(
                range(len(dates)),
                amounts,
                color="#6366f1",
                linewidth=2,
                marker="o",
                markersize=5,
            )
            step = max(1, ceil(len(dates) / 10))
            tick_positions = list(range(0, len(dates), step))
            axis.set_xticks(tick_positions)
            axis.set_xticklabels(
                [dates[position] for position in tick_positions],
                rotation=30,
                ha="right",
            )
            if len(dates) <= 20:
                for position, amount in enumerate(amounts):
                    axis.annotate(
                        f"{amount:,.2f} 元",
                        (position, amount),
                        xytext=(0, 8),
                        textcoords="offset points",
                        ha="center",
                        fontsize=8,
                    )
            else:
                peak = max(range(len(amounts)), key=amounts.__getitem__)
                axis.annotate(
                    f"{dates[peak]}：{amounts[peak]:,.2f} 元",
                    (peak, amounts[peak]),
                    xytext=(8, 10),
                    textcoords="offset points",
                    fontsize=8,
                )
            axis.set_ylabel("销售额（元）")
            axis.grid(axis="y", color="#e5e7eb", linewidth=0.8)
            axis.spines[["top", "right"]].set_visible(False)
        else:
            axis.text(
                0.5,
                0.5,
                "没有可解析的订单日期",
                ha="center",
                va="center",
                transform=axis.transAxes,
            )
            axis.set_axis_off()
        axis.set_title("每日销售趋势", loc="left", pad=12, fontsize=14)
        figure.tight_layout()
        return _to_png(figure)


def _build_bar_chart(labels, amounts, quantities, title, font, *, height=None):
    count = len(labels)
    figure_height = height or max(4.0, min(8.5, 2.6 + count * 0.38))
    display_labels = [str(label) for label in labels]
    rows = list(zip(display_labels, amounts, quantities))[::-1]
    labels_reversed, amounts_reversed, quantities_reversed = zip(*rows) if rows else ([], [], [])

    with plt.rc_context(_plot_style(font)):
        figure, axis = plt.subplots(figsize=(9, figure_height))
        if labels_reversed:
            bars = axis.barh(
                labels_reversed,
                amounts_reversed,
                color="#6366f1",
                height=0.62,
            )
            labels_for_bars = [
                f"{amount:,.2f} 元  |  销量：{quantity:,.0f} 件"
                for amount, quantity in zip(amounts_reversed, quantities_reversed)
            ]
            axis.bar_label(
                bars,
                labels=labels_for_bars,
                label_type="edge",
                padding=4,
                fontsize=8,
            )
            minimum = min([0.0, *amounts_reversed])
            maximum = max([0.0, *amounts_reversed])
            span = maximum - minimum or max(abs(maximum), 1.0)
            axis.set_xlim(minimum - span * 0.05, maximum + span * 0.58)
            axis.set_xlabel("销售额（元）")
            axis.grid(axis="x", color="#e5e7eb", linewidth=0.8)
            axis.set_axisbelow(True)
            axis.spines[["top", "right", "left"]].set_visible(False)
            axis.tick_params(axis="y", length=0)
        else:
            axis.text(
                0.5,
                0.5,
                "没有可用的商品销售数据",
                ha="center",
                va="center",
                transform=axis.transAxes,
            )
            axis.set_axis_off()
        axis.set_title(title, loc="left", pad=12, fontsize=14)
        figure.tight_layout()
        return _to_png(figure)


def _build_top_products_chart(result, font):
    trace = _chart_data(result, "product_top10")
    labels = _as_list(trace, "x")
    amounts = [_number(value) for value in _as_list(trace, "y")]
    quantities = [_number(value) for value in _as_list(trace, "customdata")]
    if len(quantities) != len(labels):
        quantities = [0.0] * len(labels)
    return _build_bar_chart(
        labels,
        amounts,
        quantities,
        "商品 Top10 销售额与销量",
        font,
    )


def _build_product_share_chart(result, font):
    trace = _chart_data(result, "product_share")
    labels = [str(value) for value in _as_list(trace, "labels")]
    amounts = [_number(value) for value in _as_list(trace, "values")]
    quantities = [_number(value) for value in _as_list(trace, "customdata")]
    if len(quantities) != len(labels):
        quantities = [0.0] * len(labels)

    # A pie chart implies nonnegative parts of a whole. Keep all sales values
    # visible in a bar chart when returns or other negative amounts are present.
    if any(value < 0 for value in amounts):
        return _build_bar_chart(
            labels,
            amounts,
            quantities,
            "商品销售额（含负数，使用柱状图展示）",
            font,
            height=max(4.5, min(10.0, 2.8 + len(labels) * 0.3)),
        )

    with plt.rc_context(_plot_style(font)):
        figure, axis = plt.subplots(figsize=(9, 5.5))
        total = sum(amounts)
        if labels and total > 0:
            wedges, _texts, _autotexts = axis.pie(
                amounts,
                startangle=90,
                counterclock=False,
                colors=plt.get_cmap("tab20").colors[: len(labels)],
                autopct=lambda percentage: f"{percentage:.1f}%" if percentage >= 4 else "",
                pctdistance=0.72,
                textprops={"fontsize": 8},
            )
            legend_labels = []
            for label, amount, quantity in zip(labels, amounts, quantities):
                share = amount / total * 100
                legend_labels.append(
                    f"{label}（销量 {quantity:,.0f} 件，{amount:,.2f} 元，{share:.1f}%）"
                )
            axis.legend(
                wedges,
                legend_labels,
                title="商品 / 销量 / 销售额 / 占比",
                loc="center left",
                bbox_to_anchor=(1.0, 0.5),
                frameon=False,
                fontsize=8,
                title_fontsize=9,
            )
            axis.axis("equal")
        else:
            axis.text(
                0.5,
                0.5,
                "销售总额不为正数，无法计算有效销售占比",
                ha="center",
                va="center",
                transform=axis.transAxes,
            )
            axis.set_axis_off()
        axis.set_title("商品销售占比", loc="left", pad=12, fontsize=14)
        figure.tight_layout()
        return _to_png(figure)


def _set_run_font(run, *, size=None, bold=None, color=None):
    run.font.name = "Microsoft YaHei"
    run_properties = run._element.get_or_add_rPr()
    run_fonts = run_properties.rFonts
    if run_fonts is None:
        run_fonts = OxmlElement("w:rFonts")
        run_properties.insert(0, run_fonts)
    run_fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def _shade_cell(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _add_table(document, headers, rows):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = str(header)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _shade_cell(cell, "365F91")
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                _set_run_font(run, size=9, bold=True, color=(255, 255, 255))
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = str(value)
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cells[index].paragraphs:
                for run in paragraph.runs:
                    _set_run_font(run, size=9, color=(31, 41, 55))
    document.add_paragraph()
    return table


def _add_heading(document, text, level=1):
    paragraph = document.add_heading(str(text), level=level)
    paragraph.paragraph_format.keep_with_next = True
    for run in paragraph.runs:
        _set_run_font(
            run,
            size=16 if level == 1 else 12,
            bold=True,
            color=(31, 78, 121),
        )
    return paragraph


def _add_paragraph(document, text, *, style=None, bold_prefix=None):
    paragraph = document.add_paragraph(style=style)
    if bold_prefix and str(text).startswith(bold_prefix):
        prefix = paragraph.add_run(bold_prefix)
        _set_run_font(prefix, bold=True, color=(31, 41, 55))
        body = paragraph.add_run(str(text)[len(bold_prefix):])
        _set_run_font(body, color=(31, 41, 55))
    else:
        run = paragraph.add_run(str(text))
        _set_run_font(run, color=(31, 41, 55))
    paragraph.paragraph_format.space_after = Pt(6)
    return paragraph


def _add_chart(document, image_bytes, alt_text):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    picture = run.add_picture(BytesIO(image_bytes), width=Inches(6.65))
    picture._inline.docPr.set("descr", alt_text)


def _top_product_rows(result):
    trace = _chart_data(result, "product_top10")
    labels = _as_list(trace, "x")
    amounts = [_number(value) for value in _as_list(trace, "y")]
    quantities = [_number(value) for value in _as_list(trace, "customdata")]
    if len(quantities) != len(labels):
        quantities = [0.0] * len(labels)
    total_sales = _number(result.get("total_sales", 0))
    return [
        [
            label,
            f"{amount:,.2f}",
            f"{quantity:,.0f}",
            f"{amount / total_sales * 100:.2f}%" if total_sales else "—",
        ]
        for label, amount, quantity in zip(labels, amounts, quantities)
    ]


def _daily_summary_rows(result):
    trace = _chart_data(result, "sales_trend")
    dates = _as_list(trace, "x")
    amounts = [_number(value) for value in _as_list(trace, "y")]
    counts = _as_list(trace, "customdata")
    if len(counts) != len(dates):
        counts = [0] * len(dates)
    ranked = sorted(
        zip(dates, amounts, counts),
        key=lambda item: item[1],
        reverse=True,
    )[:10]
    if not ranked:
        return [["无可解析日期", "—", "—"]]
    return [
        [date, f"{amount:,.2f}", f"{int(_number(count)):,}"]
        for date, amount, count in ranked
    ]


def _product_conclusion(result):
    trace = _chart_data(result, "product_top10")
    labels = _as_list(trace, "x")
    amounts = [_number(value) for value in _as_list(trace, "y")]
    if not labels or not amounts:
        return "商品字段没有可用于汇总的销售数据。"

    top_index = max(range(len(amounts)), key=amounts.__getitem__)
    top_product = labels[top_index]
    top_amount = amounts[top_index]
    total_sales = _number(result.get("total_sales", 0))
    if total_sales > 0:
        share = top_amount / total_sales * 100
        threshold_note = "超过50%的主要收入来源阈值" if share > 50 else "未超过50%的主要收入来源阈值"
        return (
            f"销售额排名第一的商品为{top_product}，销售额 {top_amount:,.2f} 元，"
            f"占总销售额 {share:.1f}%，{threshold_note}。"
        )
    return (
        f"销售额排名第一的商品为{top_product}，净销售额 {top_amount:,.2f} 元；"
        "总销售额不为正数，销售占比不作为有效构成比例解读。"
    )


def _configure_document(document):
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    for style_name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = document.styles[style_name]
        style.font.name = "Microsoft YaHei"
        properties = style._element.get_or_add_rPr()
        fonts = properties.rFonts
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            properties.insert(0, fonts)
        fonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    document.styles["Normal"].font.size = Pt(10)
    document.styles["Title"].font.size = Pt(24)
    document.styles["Heading 1"].font.size = Pt(16)
    document.styles["Heading 2"].font.size = Pt(12)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.add_run("Excel Insight  |  销售数据分析报告")
    _set_run_font(header_run, size=8, color=(107, 114, 128))
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("基于本次上传数据自动生成")
    _set_run_font(footer_run, size=8, color=(107, 114, 128))


def build_analysis_report(result):
    """Return a Word report containing page-consistent KPIs, conclusions and charts."""
    document = Document()
    _configure_document(document)
    document.core_properties.title = "Excel Insight 销售数据分析报告"
    document.core_properties.subject = "销售表现、商品分析与数据质量"

    title = document.add_paragraph()
    title.style = document.styles["Title"]
    title_run = title.add_run("销售数据分析报告")
    _set_run_font(title_run, size=24, bold=True, color=(31, 78, 121))
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
    subtitle_run = subtitle.add_run(
        f"Excel Insight 自动生成  ·  {datetime.now().astimezone().strftime('%Y-%m-%d %H:%M')}"
    )
    _set_run_font(subtitle_run, size=9, color=(107, 114, 128))

    orders = int(result.get("orders", 0))
    row_count = int(result.get("row_count", 0))
    total_sales = _number(result.get("total_sales", 0))
    average_order_sales = _number(result.get("average_order_sales", 0))
    empty_order_ids = int(result.get("empty_order_ids", 0))
    quality = result["data_quality"]
    insights = result.get("insights", {})
    sales_insights = insights.get("sales_insights", []) or []
    product_insights = insights.get("product_insights", []) or []
    data_warnings = insights.get("data_warnings", []) or []

    _add_heading(document, "分析结论")
    _add_paragraph(
        document,
        f"本次分析覆盖 {row_count:,} 行销售明细和 {orders:,} 个有效唯一订单，"
        f"总销售额为 {total_sales:,.2f} 元，平均订单金额为 {average_order_sales:,.2f} 元。"
        "订单数按非空订单编号去重；所有商品明细行均参与销售额汇总。",
    )
    for insight in sales_insights:
        _add_paragraph(document, f"销售表现：{insight}", bold_prefix="销售表现：")
    if not sales_insights:
        _add_paragraph(document, "销售表现：订单日期无可解析值，无法定位每日销售峰值。", bold_prefix="销售表现：")

    product_finding = (
        product_insights[0]
        if product_insights
        else _product_conclusion(result)
    )
    _add_paragraph(document, f"商品分析：{product_finding}", bold_prefix="商品分析：")
    for insight in product_insights[1:]:
        _add_paragraph(document, insight)
    if data_warnings:
        for warning in data_warnings:
            _add_paragraph(document, f"数据提醒：{warning}", bold_prefix="数据提醒：")
    elif not empty_order_ids:
        _add_paragraph(document, "数据提醒：未检测到缺失单元格或完整重复行。", bold_prefix="数据提醒：")
    if empty_order_ids:
        _add_paragraph(
            document,
            f"数据提醒：有 {empty_order_ids:,} 行订单编号为空或仅含空格；这些行不计入订单数，但其销售额仍计入总额。",
            bold_prefix="数据提醒：",
        )

    _add_heading(document, "核心指标", level=1)
    _add_table(
        document,
        ["指标", "结果", "统计口径"],
        [
            ["唯一订单数量", f"{orders:,} 笔", "去除首尾空格后，对非空订单编号去重"],
            ["销售明细行数", f"{row_count:,} 行", "每行按一条商品销售明细统计"],
            ["总销售额", f"{total_sales:,.2f} 元", "所有明细行销售额之和；保留重复行与负数行"],
            ["平均订单金额", f"{average_order_sales:,.2f} 元", "总销售额 ÷ 唯一订单数；订单数为0时显示0.00元"],
        ],
    )

    _add_heading(document, "销售表现", level=1)
    trend_dates = _as_list(_chart_data(result, "sales_trend"), "x")
    if trend_dates:
        _add_paragraph(
            document,
            f"图表按可解析日期汇总；展示范围为 {min(trend_dates)} 至 {max(trend_dates)}。"
            "无法解析的日期不进入趋势图，但对应明细销售额仍计入总额。",
        )
    else:
        _add_paragraph(document, "当前数据没有可解析的订单日期，趋势图将显示无可用日期提示。")
    font = _chart_font()
    _add_chart(
        document,
        _build_trend_chart(result, font),
        "按日期汇总的销售额折线图",
    )
    _add_heading(document, "每日销售额 Top10 日期", level=2)
    _add_table(
        document,
        ["日期", "销售额（元）", "非空唯一订单数"],
        _daily_summary_rows(result),
    )

    _add_heading(document, "商品表现", level=1)
    _add_paragraph(document, "下图展示销售额排名前列商品及销量，表格列出对应金额和占比。")
    _add_chart(
        document,
        _build_top_products_chart(result, font),
        "商品销售额 Top10 横向柱状图，标注销售额和销量",
    )
    _add_heading(document, "商品销售额 Top10 汇总", level=2)
    _add_table(
        document,
        ["商品", "销售额（元）", "销量（件）", "占总销售额"],
        _top_product_rows(result) or [["暂无商品数据", "—", "—", "—"]],
    )
    if any(_number(value) < 0 for value in _as_list(_chart_data(result, "product_share"), "values")):
        _add_paragraph(
            document,
            "商品明细中含有负数销售额，因此销售占比图改用商品销售额柱状图展示；负数原值保留，避免将负值误画为饼图占比。",
        )
    _add_chart(
        document,
        _build_product_share_chart(result, font),
        "商品销售结构图，商品图例包含销量、销售额与占比；负数数据以柱状图展示",
    )

    _add_heading(document, "数据质量", level=1)
    _add_table(
        document,
        ["质量指标", "结果", "说明"],
        [
            ["数据行数", f"{int(quality['rows']):,}", "上传表格中的数据行数"],
            ["字段数量", f"{int(quality['columns']):,}", "用于质量分析的原始字段数"],
            ["缺失值", f"{int(quality['missing_values']):,}", "按 Pandas 缺失值统计；空白订单编号另行统计"],
            ["完整重复行", f"{int(quality['duplicate_rows']):,}", "完全相同的整行记录；仅提示，不自动删除"],
            ["空订单编号", f"{empty_order_ids:,}", "不计入唯一订单数；对应销售明细仍计入销售额"],
            ["数据完整率", f"{float(quality['completeness']) * 100:.2f}%", "1 - 缺失单元格数 ÷ 总单元格数"],
            ["质量评分", f"{float(quality['quality_score']):.2f} / 100", "从100分扣除缺失率和重复行率对应分值，最低为0"],
        ],
    )
    _add_paragraph(
        document,
        "口径说明：缺失单元格和完整重复行只用于质量提示，不会自动填补或删除。销售额仍按数量 × 单价逐行汇总；退货等负数记录保留在总额和商品汇总中。",
    )

    output = BytesIO()
    document.save(output)
    return output.getvalue()
