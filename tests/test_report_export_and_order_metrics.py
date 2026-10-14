import io
import re

import pandas as pd
from docx import Document

from app import app


def _workbook_bytes(dataframe):
    output = io.BytesIO()
    dataframe.to_excel(output, index=False)
    output.seek(0)
    return output


def test_word_report_matches_page_metrics_and_contains_conclusions_tables_and_charts():
    sales = pd.DataFrame({
        "order_id": [
            "ORD001",
            " ORD001 ",
            "ORD001",
            "ORD002",
            None,
            "   ",
            "ORD003",
        ],
        "order_date": [
            "2026-10-01",
            "2026-10-01",
            "2026-10-01",
            "2026-10-02",
            "2026-10-03",
            "2026-10-03",
            "2026-10-04",
        ],
        "product": ["手机", "耳机", "手机", "键盘", "显示器", "鼠标", "退货"],
        "quantity": [1, 2, 1, 1, 1, 1, 1],
        "unit_price": [100, 50, 100, 200, 1000, 500, -30],
    })
    upload = _workbook_bytes(sales)

    with app.test_client() as client:
        page = client.post(
            "/",
            data={"file": (upload, "multi-line-orders.xlsx")},
            content_type="multipart/form-data",
        )
        html = page.get_data(as_text=True)
        assert page.status_code == 200
        assert "订单数量：3" in html
        assert "1970.00 元" in html
        assert "656.67 元" in html
        assert "空编号 2 行不计入订单数" in html
        assert "发现1条重复数据" in html
        assert "下载 Word 分析报告" in html

        report_link = re.search(r'href="(/download-report/[^\"]+)"', html)
        assert report_link is not None
        download = client.get(report_link.group(1))

    assert download.status_code == 200
    assert download.mimetype == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert download.headers["Content-Disposition"].endswith(
        'filename=excel_insight_analysis_report.docx'
    )

    report = Document(io.BytesIO(download.data))
    assert len(report.tables) >= 4
    assert len(report.inline_shapes) == 3

    kpis = report.tables[0]
    assert kpis.cell(1, 1).text == "3 笔"
    assert kpis.cell(2, 1).text == "7 行"
    assert kpis.cell(3, 1).text == "1,970.00 元"
    assert kpis.cell(4, 1).text == "656.67 元"

    daily_summary = report.tables[1]
    daily_order_counts = {
        daily_summary.cell(row, 0).text: daily_summary.cell(row, 2).text
        for row in range(1, 5)
    }
    assert daily_order_counts == {
        "2026-10-01": "1",
        "2026-10-02": "1",
        "2026-10-03": "0",
        "2026-10-04": "1",
    }

    product_summary = report.tables[2]
    assert product_summary.cell(1, 0).text == "显示器"
    assert product_summary.cell(1, 1).text == "1,000.00"
    assert product_summary.cell(1, 2).text == "1"

    quality = report.tables[3]
    assert quality.cell(3, 1).text == "1"
    assert quality.cell(4, 1).text == "1"
    assert quality.cell(5, 1).text == "2"
    assert "不计入唯一订单数" in quality.cell(5, 2).text
    assert "不自动删除" in quality.cell(4, 2).text

    all_report_text = " ".join(
        [paragraph.text for paragraph in report.paragraphs]
        + [cell.text for table in report.tables for row in table.rows for cell in row.cells]
    )
    assert "2026-10-03销售额最高" in all_report_text
    assert "显示器贡献50.8%的销售额" in all_report_text
    assert "含有负数销售额" in all_report_text
    assert "C:/private/" not in all_report_text


def test_report_is_not_offered_when_internal_error_occurs(monkeypatch):
    def fail_loader(_file):
        raise ValueError("C:/private/internal.xlsx")

    monkeypatch.setattr("app.load_sales_data", fail_loader)
    upload = _workbook_bytes(pd.DataFrame({"sample": [1]}))

    with app.test_client() as client:
        response = client.post(
            "/",
            data={"file": (upload, "internal-error.xlsx")},
            content_type="multipart/form-data",
        )

    html = response.get_data(as_text=True)
    assert "文件无法处理，请确认Excel格式正确" in html
    assert "C:/private/internal.xlsx" not in html
    assert "下载 Word 分析报告" not in html


def test_blank_order_ids_do_not_divide_average_by_zero():
    sales = pd.DataFrame({
        "order_id": [None, "   "],
        "order_date": ["2026-10-01", "2026-10-01"],
        "product": ["样例商品", "样例商品"],
        "quantity": [1, 2],
        "unit_price": [100, 50],
    })
    upload = _workbook_bytes(sales)

    with app.test_client() as client:
        page = client.post(
            "/",
            data={"file": (upload, "blank-order-ids.xlsx")},
            content_type="multipart/form-data",
        )
        html = page.get_data(as_text=True)
        assert "订单数量：0" in html
        assert "200.00 元" in html
        assert "0.00 元" in html
        report_link = re.search(r'href="(/download-report/[^\"]+)"', html)
        assert report_link is not None
        download = client.get(report_link.group(1))

    report = Document(io.BytesIO(download.data))
    assert report.tables[0].cell(1, 1).text == "0 笔"
    assert report.tables[0].cell(3, 1).text == "200.00 元"
    assert report.tables[0].cell(4, 1).text == "0.00 元"
    assert report.tables[1].cell(1, 2).text == "0"
    assert report.tables[3].cell(5, 1).text == "2"
